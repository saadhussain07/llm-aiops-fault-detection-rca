#!/usr/bin/env python3
"""
rca_blind_sheet.py - build the blinded rating sheet for the secondary RCA
check (FGCS v2 protocol, Amendment 2, section 1).

For every condition, the root-cause narrative of the first
ANOMALY_CONFIRMED cycle of each run (S1-S5) is collected, exactly as
rca_metrics.py selects it. Identical narratives are rated once. The
narratives are shuffled with a fixed seed and written to one sheet per
rater; the condition, model, run and cycle are hidden. The key, which
maps each item back to its conditions and holds the mechanical scores,
is written separately and must not be opened by the raters.

Usage:
  python analysis/rca_blind_sheet.py <data-dir> live replay/<dir> ...

Writes rca_rating_A.xlsx, rca_rating_B.xlsx and rca_rating_key.csv
(requires openpyxl).
"""

import csv
import random
import sys
from pathlib import Path

from rca_metrics import first_class, first_service, load, norm, score

SEED = 20260929
FAULT = {
    "S1_CPU_STRESS": "CPU stress: busy-loop started inside cartservice",
    "S2_POD_CRASH": "Pod crash: cartservice pod force-deleted",
    "S3_MEMORY_PRESSURE": "Memory pressure: heap allocation inside paymentservice",
    "S4_NETWORK_LATENCY": "Network latency: 500 ms delay added with tc netem on frontend",
    "S5_CASCADING_FAILURE": "Dependency failure: redis-cart (the cart store) force-deleted",
}
CLASS_TEXT = {
    "resource": "resource exhaustion in the service itself (CPU or memory)",
    "crash": "the process or pod terminated and restarted",
    "network": "added network delay / latency",
    "dependency": "failure of a backing dependency the service relies on",
}
INSTRUCTIONS = [
    "Rate every row. You do not know which system wrote a narrative; do not try to find out.",
    "Do not discuss items with the other rater until both sheets are complete.",
    "service_ok = 1 if the narrative names the injected service as the ROOT CAUSE "
    "(not merely as a service that shows symptoms), else 0.",
    "class_ok = 1 if the narrative attributes the problem to the injected fault class "
    "as its cause (see the class column), else 0. A class mentioned only as a "
    "symptom or a downstream effect does not count.",
    "Judge only the text shown. Use the notes column for doubtful cases.",
]


def first_confirmed(live, rep):
    out = {}
    runs = {}
    for k in live:
        runs.setdefault((k[0], k[1]), []).append(k)
    for (scen, run), ks in runs.items():
        if scen not in FAULT:
            continue
        for k in sorted(ks):
            v = rep[k]["llm_verdict"] if k in rep else live[k].get("llm_verdict")
            if v and v.get("verdict") == "ANOMALY_CONFIRMED":
                out[(scen, run)] = (k, v)
                break
    return out


def main(data_dir, specs):
    live = load(data_dir)
    items = {}                                  # narrative -> item
    for spec in specs:
        rep = {} if spec == "live" else load(spec)
        name = "live" if spec == "live" else Path(spec).name
        for (scen, run), (k, v) in sorted(first_confirmed(live, rep).items()):
            text = norm(v.get("root_cause", "")).strip()
            ok, s, c = score(live[k], v)
            it = items.setdefault((scen, text), {"scen": scen, "text": text,
                                                 "uses": [], "mech": ok,
                                                 "svc": s, "cls": c})
            it["uses"].append(f"{name}|run{run}|cycle{k[2]}")

    order = sorted(items)
    random.Random(SEED).shuffle(order)
    rows = []
    for i, key in enumerate(order, 1):
        it = items[key]
        it["id"] = f"N{i:03d}"
        rows.append(it)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    for rater in ("A", "B"):
        wb = Workbook()
        ws = wb.active
        ws.title = "rate"
        hdr = ["id", "injected fault", "expected root-cause service",
               "expected fault class", "narrative (root_cause)",
               "service_ok (1/0)", "class_ok (1/0)", "notes"]
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
        for it in rows:
            live0 = next(k for k in live if k[0] == it["scen"])
            exp_s = live[live0]["expected_root_cause_service"]
            exp_c = live[live0]["expected_fault_class"]
            ws.append([it["id"], FAULT[it["scen"]], exp_s,
                       f"{exp_c}: {CLASS_TEXT[exp_c]}", it["text"], None, None, None])
        for col, w in zip("ABCDEFGH", (7, 30, 16, 30, 90, 10, 10, 30)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F2:G{len(rows) + 1}")
        ws.freeze_panes = "B2"
        info = wb.create_sheet("instructions", 0)
        info.column_dimensions["A"].width = 120
        info.append([f"Rater {rater}: blinded root-cause rating ({len(rows)} items)"])
        info["A1"].font = Font(bold=True)
        for line in INSTRUCTIONS:
            info.append([line])
        for row in info.iter_rows():
            row[0].alignment = Alignment(wrap_text=True)
        wb.save(f"rca_rating_{rater}.xlsx")

    with open("rca_rating_key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "scenario", "uses", "mech_service", "mech_class",
                    "rule_service", "rule_class"])
        for it in rows:
            w.writerow([it["id"], it["scen"], ";".join(it["uses"]),
                        int(it["mech"]["service"]), int(it["mech"]["class"]),
                        it["svc"], it["cls"]])
    n_uses = sum(len(it["uses"]) for it in rows)
    print(f"{len(rows)} unique narratives ({n_uses} condition-run slots) -> "
          f"rca_rating_A.xlsx, rca_rating_B.xlsx, rca_rating_key.csv")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2:])
