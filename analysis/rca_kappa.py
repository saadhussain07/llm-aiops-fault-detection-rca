#!/usr/bin/env python3
"""
rca_kappa.py - score the blinded RCA rating (Amendment 2, section 1).

Reports Cohen's kappa between the two raters (service, class), each
rater's agreement with the mechanical rule, and the per-condition
first-confirmed-cycle counts under each rater and under the rule.
Runs without a confirmed cycle count as misses (denominator 5 per
scenario), as in rca_metrics.py.

Usage:
  python analysis/rca_kappa.py rca_rating_A.xlsx [rca_rating_B.xlsx] rca_rating_key.csv

With one sheet, only rater-versus-rules agreement is reported.
"""

import csv
import sys
from collections import defaultdict

SCENS = ["S1_CPU_STRESS", "S2_POD_CRASH", "S3_MEMORY_PRESSURE",
         "S4_NETWORK_LATENCY", "S5_CASCADING_FAILURE"]
RUNS = 5


def read_sheet(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["rate"]
    out = {}
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if not row or not row[0]:
            continue
        s, c = row[5], row[6]
        if s in (None, "") or c in (None, ""):
            sys.exit(f"{path}: item {row[0]} is not fully rated")
        out[row[0]] = (int(s), int(c))
    return out


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0, po


def main(fa, fb, fkey):
    A = read_sheet(fa)
    B = read_sheet(fb) if fb else None
    key = list(csv.DictReader(open(fkey, encoding="utf-8")))
    ids = [k["id"] for k in key]
    missing = [i for i in ids if i not in A or (B is not None and i not in B)]
    if missing:
        sys.exit(f"Unrated items: {missing[:10]}")
    mech = {k["id"]: (int(k["mech_service"]), int(k["mech_class"])) for k in key}

    print(f"\n{len(ids)} items\n")
    print(f"{'':<26}{'service':>18}{'class':>18}")
    pairs = [("rater A vs rule", A, mech)]
    if B is not None:
        pairs = [("rater A vs rater B", A, B)] + pairs + [("rater B vs rule", B, mech)]
    pairs.append(("rater A vs rule, both", {i: (A[i][0] * A[i][1],) * 2 for i in ids},
                  {i: (mech[i][0] * mech[i][1],) * 2 for i in ids}))
    for name, X, Y in pairs:
        cells = []
        for j in (0, 1):
            k, po = kappa([X[i][j] for i in ids], [Y[i][j] for i in ids])
            cells.append(f"k={k:.2f} ({100 * po:.0f}%)")
        print(f"{name:<26}{cells[0]:>18}{cells[1]:>18}")

    # per condition: both-correct on the first confirmed cycle of each run
    counts = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    for k in key:
        for use in k["uses"].split(";"):
            cond = use.split("|")[0]
            for name, X in (("A", A), ("B", B), ("rule", mech)):
                if X is None:
                    continue
                s, c = X[k["id"]]
                counts[cond][name][k["scenario"]] += int(s and c)
    print("\nDisagreements with the rules (id, scenario, rater, rule, uses)")
    for k in key:
        if A[k["id"]] != mech[k["id"]]:
            print(f"   {k['id']} {k['scenario'][:2]} rater={A[k['id']]} rule={mech[k['id']]} {k['uses']}")
    print("\nBoth correct, first confirmed cycle per run (misses count against)")
    print(f"{'condition':<36}{'rater':<7}" + "".join(f"{s[:2]:>6}" for s in SCENS) + f"{'total':>9}")
    for cond in sorted(counts):
        for name in ("A", "B", "rule"):
            if name == "B" and B is None:
                continue
            row = counts[cond][name]
            tot = sum(row[s] for s in SCENS)
            print(f"{cond:<36}{name:<7}" + "".join(f"{row[s]:>4}/5" for s in SCENS)
                  + f"{tot:>6}/{RUNS * len(SCENS)}")


if __name__ == "__main__":
    if len(sys.argv) == 4:
        main(*sys.argv[1:])
    elif len(sys.argv) == 3:
        main(sys.argv[1], None, sys.argv[2])
    else:
        sys.exit(__doc__)
